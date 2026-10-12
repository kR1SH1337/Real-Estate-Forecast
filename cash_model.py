#!/usr/bin/env python3
"""Moscow cash/installment annual sales and area: Huber in Jan-2023 prices.

Primary model uses all available history. Rolling is an optional comparison.
No pre-smoothing or automatic imputation.
Run `python cash_model.py --help`. Units: sales in RUB billion, rates in percent.
"""
from __future__ import annotations

import argparse
from io import BytesIO
import json
from pathlib import Path

import numpy as np
import pandas as pd
import openpyxl
from scipy.optimize import minimize
from scipy.stats import chi2, shapiro

HUBER_C = 1.345
MAD_NORMAL = 1.4826022185
MONTHS_RU = ['Январь', 'Февраль', 'Март', 'Апрель', 'Май', 'Июнь',
             'Июль', 'Август', 'Сентябрь', 'Октябрь', 'Ноябрь', 'Декабрь']


def monthly_series(values: dict, label: str) -> pd.Series:
    result = pd.Series(values, dtype=float).sort_index()
    if result.empty or result.index.has_duplicates:
        raise ValueError(f'{label}: empty series or duplicate months')
    if not all(d.day == 1 for d in result.index):
        raise ValueError(f'{label}: dates must identify the first day of a month')
    return result


def read_sources(cash_path: Path, macro_path: Path, price_path: Path,
                 as_of: str | None = None) -> tuple[pd.DataFrame, pd.Series]:
    """Read actual cash row, Moscow KS and Moscow/general/all-rooms price index."""
    sheet = openpyxl.load_workbook(cash_path, data_only=True).worksheets[0]
    values = {}
    for col in range(2, sheet.max_column + 1):
        date = sheet.cell(1, col).value
        if date is None:
            continue
        date = pd.Timestamp(date)
        if date in values:
            raise ValueError('Duplicate month in cash workbook')
        value = sheet.cell(2, col).value
        values[date] = np.nan if value is None else float(value)
    cash = monthly_series(values, 'cash')
    cash = cash.reindex(pd.date_range(cash.index.min(), cash.index.max(), freq='MS'))
    if as_of:
        cutoff = pd.Timestamp(as_of)
        if cutoff.day != 1:
            raise ValueError('--as-of must be the first day of a month')
        cash = cash.loc[:cutoff]
    if cash.empty or not np.isfinite(cash.iloc[-1]):
        raise ValueError('The last cash observation must exist')

    sheet = openpyxl.load_workbook(macro_path, data_only=True)['Москва']
    values = {}
    for row in range(2, sheet.max_row + 1):
        date, value = sheet.cell(row, 1).value, sheet.cell(row, 2).value
        if date is None:
            continue
        date = pd.Timestamp(date)
        if date in values:
            raise ValueError('Duplicate month in KS workbook')
        values[date] = np.nan if value is None else float(value)
    rates = monthly_series(values, 'KS')
    rates = rates.reindex(pd.date_range(rates.index.min(), rates.index.max(), freq='MS'))

    sheet = openpyxl.load_workbook(price_path, data_only=True).worksheets[0]
    rows = [r for r in range(2, sheet.max_row + 1)
            if [sheet.cell(r, c).value for c in (1, 2, 3)]
            == ['Город Москва', 'Общий', 'Все']]
    if len(rows) != 1:
        raise ValueError('Expected exactly one Moscow / Общий / Все price row')
    values = {}
    for col in range(4, sheet.max_column + 1):
        header = sheet.cell(1, col).value
        if header is None:
            continue
        month, year = str(header).split()
        date = pd.Timestamp(int(year), MONTHS_RU.index(month) + 1, 1)
        if date in values:
            raise ValueError('Duplicate month in price index')
        values[date] = float(sheet.cell(rows[0], col).value)
    prices = monthly_series(values, 'price index')
    base = prices.loc[pd.Timestamp('2023-01-01')]
    if not np.isfinite(base) or base <= 0:
        raise ValueError('Invalid January-2023 price-index base')
    data = pd.DataFrame(index=cash.index)
    data.index.name = 'date'
    data['nominal_source'] = cash
    # Non-positive reconstructed residuals are censored/missing, not true zero sales.
    data['nominal'] = cash.mask(cash <= 0)
    data['price_factor'] = prices.reindex(data.index) / base
    data['ks_avg3'] = rates.shift(1).rolling(3, min_periods=3).mean().reindex(data.index)
    if not np.isfinite(data[['price_factor', 'ks_avg3']].to_numpy()).all():
        raise ValueError('Missing prices or previous three KS observations; no imputation')
    if (data['price_factor'] <= 0).any():
        raise ValueError('Price-index values must be positive')
    data['real'] = data['nominal'] / data['price_factor']
    return data, rates


def fit_model(y, x, method='huber') -> dict:
    """OLS, Huber IRLS with iterated residual MAD, or Student-t MLE with fixed df=4."""
    y, x = np.asarray(y, float), np.asarray(x, float)
    valid = np.isfinite(y) & np.isfinite(x)
    yy, xx = y[valid], x[valid]
    if len(yy) < 6 or np.ptp(xx) < 1e-8:
        raise ValueError('At least 6 usable observations and varying KS required')
    # Numerical centering/scaling only. Returned coefficients use untransformed KS.
    X = np.column_stack([np.ones(len(xx)), (xx - 16) / 5])
    b = np.linalg.lstsq(X, yy, rcond=None)[0]
    if method == 'ols':
        scale = float(np.std(yy - X @ b, ddof=2))
        weights, iterations = np.ones(len(yy)), 0
    elif method == 'huber':
        for iterations in range(1000):
            e = yy - X @ b
            scale = max(MAD_NORMAL * np.median(abs(e - np.median(e))), 1e-6)
            weights = np.minimum(1, HUBER_C / np.maximum(abs(e / scale), 1e-15))
            root_w = np.sqrt(weights)
            candidate = np.linalg.lstsq(X * root_w[:, None], yy * root_w, rcond=None)[0]
            if np.linalg.norm(candidate - b) < 1e-9 * (1 + np.linalg.norm(b)):
                b = candidate
                break
            b = candidate
        else:
            raise RuntimeError('Huber IRLS did not converge')
        e = yy - X @ b
        scale = max(MAD_NORMAL * np.median(abs(e - np.median(e))), 1e-6)
        weights = np.minimum(1, HUBER_C / np.maximum(abs(e / scale), 1e-15))
    elif method == 'student4':
        nu = 4.0
        residuals = yy - X @ b
        initial_scale = max(1.4826 * np.median(abs(residuals - np.median(residuals))), 1.0)

        def objective(parameters):
            e = yy - X @ parameters[:2]
            s = np.exp(parameters[2])
            u = e / s
            w = (nu + 1) / (nu + u * u)
            value = len(yy) * parameters[2] + (nu + 1) / 2 * np.log1p(u * u / nu).sum()
            gradient = np.r_[-X.T @ (w * e) / s**2, len(yy) - sum(w * u * u)]
            return value, gradient

        candidates = [minimize(objective, np.r_[start, np.log(initial_scale)], jac=True,
                               method='L-BFGS-B', bounds=[(None, None), (None, None), (-5, 8)],
                               options={'ftol': 1e-13, 'gtol': 1e-7, 'maxiter': 1000})
                      for start in [b, np.array([np.median(yy), 0.0])]]
        best = min(candidates, key=lambda item: item.fun)
        if not best.success and np.linalg.norm(best.jac) > 1e-4:
            raise RuntimeError(f'Student-t fit failed: {best.message}')
        b, scale, iterations = best.x[:2], float(np.exp(best.x[2])), best.nit
        weights = (nu + 1) / (nu + ((yy - X @ b) / scale)**2)
    else:
        raise ValueError(f'Unknown method: {method}')
    coefficients = np.array([b[0] - 16 * b[1] / 5, b[1] / 5])
    return {'coef': coefficients, 'weights': weights, 'scale': float(scale),
            'valid': valid, 'residuals': yy - coefficients[0] - coefficients[1] * xx,
            'iterations': int(iterations), 'method': method}


def predict(model: dict, x) -> np.ndarray:
    return model['coef'][0] + model['coef'][1] * np.asarray(x, float)


def metrics(actual, predicted) -> dict:
    actual, predicted = np.asarray(actual, float), np.asarray(predicted, float)
    if len(actual) == 0:
        return {'n': 0}
    if not np.isfinite(actual).all() or not np.isfinite(predicted).all() or (actual <= 0).any():
        raise ValueError('Metrics require finite predictions and positive annual actuals')
    error = predicted - actual
    return {'n': len(actual), 'rmse': float(np.sqrt(np.mean(error**2))),
            'mape_pct': float(np.mean(abs(error) / actual) * 100),
            'wape_pct': float(sum(abs(error)) / sum(actual) * 100),
            'bias_pct': float(sum(error) / sum(actual) * 100),
            'max_abs_pct': float(max(abs(error) / actual) * 100)}


def backtest(data: pd.DataFrame, window: int) -> tuple[pd.DataFrame, dict]:
    """Re-fit at each origin. Future KS/prices are known; horizons overlap."""
    records = []
    for end in range(window, len(data) - 11):
        future = data.iloc[end:end + 12]
        # A missing annual actual cannot silently disappear in a sum.
        if future['nominal'].isna().any():
            continue
        for method in ['huber', 'ols', 'student4']:
            for scope in ['rolling', 'expanding']:
                train = data.iloc[max(0, end - window) if scope == 'rolling' else 0:end]
                model = fit_model(train['real'], train['ks_avg3'], method)
                real_prediction = predict(model, future['ks_avg3'])
                nominal_prediction = real_prediction * future['price_factor'].to_numpy()
                records.append({'method': method, 'scope': scope,
                                'origin': data.index[end - 1], 'forecast_start': future.index[0],
                                'forecast_end': future.index[-1],
                                'actual_nominal': future['nominal'].sum(),
                                'predicted_nominal': nominal_prediction.sum(),
                                'actual_real': future['real'].sum(),
                                'predicted_real': real_prediction.sum(),
                                'alpha': model['coef'][0], 'beta': model['coef'][1],
                                'n_train': int(model['valid'].sum())})
    results = pd.DataFrame(records)
    summary = {}
    if not results.empty:
        for (method, scope), group in results.groupby(['method', 'scope']):
            summary[f'{method}_{scope}'] = {
                unit: metrics(group['actual_' + unit], group['predicted_' + unit])
                for unit in ['nominal', 'real']}
    return results, summary


def coefficient_history(data: pd.DataFrame, window: int, scope='expanding') -> pd.DataFrame:
    records = []
    for end in range(window, len(data) + 1):
        train = data.iloc[end - window:end] if scope == 'rolling' else data.iloc[:end]
        model = fit_model(train['real'], train['ks_avg3'])
        records.append({'date': data.index[end - 1], 'scope': scope, 'window_start': train.index[0],
                        'alpha': model['coef'][0], 'beta': model['coef'][1],
                        'level_at_16': float(predict(model, [16])[0]),
                        'n': int(model['valid'].sum()),
                        'ks_min': train['ks_avg3'].min(), 'ks_max': train['ks_avg3'].max()})
    return pd.DataFrame(records)


def diagnostics(train: pd.DataFrame, model: dict, bootstrap: int, seed: int) -> dict:
    """Descriptive residual tests; bootstrap CI is NOT a forecast interval."""
    valid = model['valid']
    e, x = model['residuals'], train['ks_avg3'].to_numpy()[valid]
    # Preserve calendar gaps: use the contiguous tail after the last missing Y.
    residuals = np.full(len(train), np.nan)
    residuals[valid] = e
    missing = np.flatnonzero(~np.isfinite(residuals))
    tail = residuals[missing[-1] + 1:] if len(missing) else residuals
    tail = tail - tail.mean()
    acf, lb = {}, {}
    for lag in range(1, min(12, len(tail) - 1) + 1):
        acf[lag] = float(tail[lag:] @ tail[:-lag] / (tail @ tail))
    for lag in [6, 12]:
        if lag in acf:
            statistic = len(tail) * (len(tail) + 2) * sum(
                acf[j]**2 / (len(tail) - j) for j in range(1, lag + 1))
            lb[lag] = {'statistic': float(statistic), 'p_value': float(chi2.sf(statistic, lag))}
    X = np.column_stack([np.ones(len(x)), x])
    q = e**2
    fitted_q = X @ np.linalg.lstsq(X, q, rcond=None)[0]
    q_variance = sum((q - q.mean())**2)
    bp_stat = len(e) * (1 - sum((q - fitted_q)**2) / q_variance) if q_variance > 0 else 0.0
    loo = []
    for position in np.flatnonzero(valid):
        yy = train['real'].to_numpy().copy()
        yy[position] = np.nan
        fitted = fit_model(yy, train['ks_avg3'], model['method'])
        loo.append({'date': str(train.index[position].date()), 'beta': float(fitted['coef'][1])})
    result = {'n': len(e), 'alpha': float(model['coef'][0]), 'beta': float(model['coef'][1]),
              'scale_mad_real': model['scale'], 'huber_threshold_real': HUBER_C * model['scale'],
              'monthly_rmse_real': float(np.sqrt(np.mean(e**2))),
              'mean_residual_real': float(e.mean()), 'acf': acf, 'ljung_box': lb,
              'koenker_bp_approx': {'statistic': float(bp_stat), 'p_value': float(chi2.sf(bp_stat, 1))},
              'shapiro_p_value': float(shapiro(e).pvalue),
              'loo_beta_range': [min(item['beta'] for item in loo), max(item['beta'] for item in loo)],
              'leave_one_out': loo, 'weights_below_one': int(sum(model['weights'] < .999)),
              'ks_range': [float(min(x)), float(max(x))],
              'notes': ['Small sample; residual p-values are approximate.',
                        'Wild bootstrap assumes symmetric zero-mean errors and fixed regressors.',
                        'Neither residual tests nor bootstrap establish parameter stability or causality.']}
    if bootstrap:
        rng = np.random.default_rng(seed)
        leverage = np.einsum('ij,jk,ik->i', X, np.linalg.inv(X.T @ X), X)
        errors = e / np.sqrt(1 - leverage)
        draws = []
        for _ in range(bootstrap):
            yy = train['real'].to_numpy().copy()
            yy[valid] = X @ model['coef'] + errors * rng.choice([-1.0, 1.0], len(e))
            draws.append(fit_model(yy, train['ks_avg3'], model['method'])['coef'])
        result['wild_bootstrap'] = {'replications': bootstrap, 'seed': seed,
                                    'ci95_alpha_beta': np.quantile(draws, [.025, .975], axis=0).T.tolist()}
    return result


def forecasts(data, rates, model, horizon, base_price_rub_m2=282700.0):
    """Annual real sales and implied area. No future price scenario is required."""
    if not np.isfinite(base_price_rub_m2) or base_price_rub_m2 <= 0:
        raise ValueError('Base price must be positive RUB per square metre')
    dates = pd.date_range(data.index[-1] + pd.offsets.MonthBegin(1), periods=horizon, freq='MS')
    factor = rates.shift(1).rolling(3, min_periods=3).mean().reindex(dates)
    if not np.isfinite(factor.to_numpy()).all():
        raise ValueError('KS scenario does not cover the forecast horizon')
    result = pd.DataFrame({'ks_avg3': factor, 'real_bln_rub': predict(model, factor)}, index=dates)
    result.index.name = 'date'
    # RUB billion / (RUB per m²) = m²; divide by one million for million m².
    result['area_m2'] = result['real_bln_rub'] * 1e9 / base_price_rub_m2
    result['area_mln_m2'] = result['area_m2'] / 1e6
    last_train_x = data.iloc[-len(model['valid']):]['ks_avg3']
    result['outside_training_ks_range'] = ((factor < last_train_x.min()) | (factor > last_train_x.max()))
    result['negative_real_forecast'] = result['real_bln_rub'] < 0
    annual = []
    for year, group in result.groupby(result.index.year):
        actual = data.loc[data.index.year == year]
        available = actual.loc[actual['real'].notna()]
        months = set(available.index.month) | set(group.index.month)
        incomplete_actuals = actual['real'].isna().any()
        actual_real = float(available['real'].sum())
        forecast_real = float(group['real_bln_rub'].sum())
        total_real = np.nan if incomplete_actuals else actual_real + forecast_real
        full_year = len(months) == 12 and not incomplete_actuals
        status = 'history_plus_forecast' if len(actual) else 'forecast'
        if not full_year:
            status += '_partial_year'
        annual.append({'year': year, 'real_bln_rub': total_real,
                       'area_mln_m2': total_real * 1e3 / base_price_rub_m2,
                       'status': status, 'actual_months': len(available),
                       'forecast_months': len(group), 'months_covered': len(months),
                       'full_year': full_year,
                       'actual_real_bln_rub': actual_real,
                       'forecast_real_bln_rub': forecast_real,
                       'actual_area_mln_m2': actual_real * 1e3 / base_price_rub_m2,
                       'forecast_area_mln_m2': forecast_real * 1e3 / base_price_rub_m2})
    return result, pd.DataFrame(annual)


def make_charts(data, train, model, history, output):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    plt.rcParams.update({'font.family': 'DejaVu Sans', 'axes.spines.top': False,
                         'axes.spines.right': False, 'figure.dpi': 140})
    fig, axes = plt.subplots(2, 1, figsize=(10, 6.4), sharex=True)
    axes[0].plot(history['date'], history['beta'], color='#155a85', marker='o', markersize=3)
    axes[0].set_ylabel('Млрд руб. на 1 п.п. КС')
    label = 'вся доступная история' if history['scope'].iloc[-1] == 'expanding' else 'rolling window'
    axes[0].set_title('Коэффициент КС — Хьюбер, ' + label)
    axes[1].plot(history['date'], history['alpha'], color='#b46b22', marker='o', markersize=3)
    axes[1].set_ylabel('Млрд руб., цены января 2023')
    axes[1].set_title('Константа')
    for ax in axes: ax.grid(alpha=.2)
    fig.autofmt_xdate(); fig.tight_layout(); buffer = BytesIO(); fig.savefig(buffer, format='png'); (output / 'coefficients.png').write_bytes(buffer.getvalue()); plt.close(fig)
    fig, axes = plt.subplots(2, 1, figsize=(10, 6.4), sharex=True)
    axes[0].plot(data.index, data['real'], color='#155a85', label='Исходный ряд в постоянных ценах')
    axes[0].plot(train.index, predict(model, train['ks_avg3']), color='#b46b22', label='Последняя оценка Хьюбера')
    axes[0].set_ylabel('Млрд руб., цены января 2023'); axes[0].legend(fontsize=8)
    dates = train.index[model['valid']]
    axes[1].bar(dates, model['residuals'], width=20, color='#155a85')
    axes[1].axhline(0, color='#333333', linewidth=.8); axes[1].set_ylabel('Остатки, млрд руб.')
    for ax in axes: ax.grid(alpha=.2)
    fig.autofmt_xdate(); fig.tight_layout(); buffer = BytesIO(); fig.savefig(buffer, format='png'); (output / 'fit_and_residuals.png').write_bytes(buffer.getvalue()); plt.close(fig)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--cash', type=Path, default=Path('data/ряд кеш.xlsx'))
    parser.add_argument('--macro', type=Path, default=Path('data/Продажи - Extended_new_07 (1).xlsx'))
    parser.add_argument('--prices', type=Path, default=Path('data/analytics_3901335506.xlsx'))
    parser.add_argument('--scope', choices=['expanding', 'rolling'], default='expanding',
                        help='Primary default: all available history; rolling is optional')
    parser.add_argument('--window', type=int, default=24,
                        help='Rolling length and first backtest origin; does not truncate expanding fit')
    parser.add_argument('--as-of', help='Last known sales month, YYYY-MM-01')
    parser.add_argument('--end-year', type=int, default=2030, help='Default forecast through December 2030')
    parser.add_argument('--horizon', type=int, help='Optional explicit months; overrides --end-year')
    parser.add_argument('--base-price-rub-m2', type=float, default=282700.0,
                        help='Actual January-2023 market price, RUB/m²; default 282700')
    parser.add_argument('--bootstrap', type=int, default=1000)
    parser.add_argument('--seed', type=int, default=20261009)
    parser.add_argument('--output', type=Path, default=Path('results'))
    parser.add_argument('--no-charts', action='store_true')
    args = parser.parse_args()
    if args.window < 12 or (args.horizon is not None and args.horizon < 1) or args.bootstrap < 0:
        parser.error('window >= 12, horizon >= 1, bootstrap >= 0 required')
    if not np.isfinite(args.base_price_rub_m2) or args.base_price_rub_m2 <= 0:
        parser.error('--base-price-rub-m2 must be positive RUB/m²')
    data, rates = read_sources(args.cash, args.macro, args.prices, args.as_of)
    horizon = args.horizon if args.horizon is not None else (
        (args.end_year - data.index[-1].year) * 12 + 12 - data.index[-1].month)
    if horizon < 1:
        parser.error('--end-year must extend beyond the last actual month')
    if len(data) < args.window:
        parser.error('Not enough calendar months for the requested window')
    train = data if args.scope == 'expanding' else data.iloc[-args.window:]
    model = fit_model(train['real'], train['ks_avg3'])
    history = coefficient_history(data, args.window, args.scope)
    tests, summary = backtest(data, args.window)
    diag = diagnostics(train, model, args.bootstrap, args.seed)
    monthly, annual = forecasts(data, rates, model, horizon, args.base_price_rub_m2)
    data['area_mln_m2'] = data['real'] * 1e3 / args.base_price_rub_m2
    if not tests.empty:
        tests['actual_area_mln_m2'] = tests['actual_real'] * 1e3 / args.base_price_rub_m2
        tests['predicted_area_mln_m2'] = tests['predicted_real'] * 1e3 / args.base_price_rub_m2
        for (method, scope), group in tests.groupby(['method', 'scope']):
            summary[f'{method}_{scope}']['area'] = metrics(group['actual_area_mln_m2'], group['predicted_area_mln_m2'])
    weights = train.loc[model['valid'], ['nominal', 'real', 'ks_avg3']].copy()
    weights['fitted_real'] = predict(model, weights['ks_avg3'])
    weights['residual_real'] = model['residuals']; weights['weight'] = model['weights']
    args.output.mkdir(parents=True, exist_ok=True)
    for filename, table, index in [('input_series.csv', data, True), ('weights.csv', weights, True),
                                  ('coefficients.csv', history, False), ('backtest.csv', tests, False),
                                  ('forecast_monthly.csv', monthly, True), ('forecast_annual.csv', annual, False)]:
        table.to_csv(args.output / filename, index=index, encoding='utf-8-sig')
    metadata = {'as_of': str(data.index[-1].date()), 'scope': args.scope, 'training_calendar_months': len(train),
                'rolling_comparison_window': args.window,
                'window_start': str(train.index[0].date()), 'price_base': '2023-01-01',
                'diagnostics': diag, 'validation': summary,
                'base_price_rub_m2': args.base_price_rub_m2,
                'forecast_end': str(monthly.index[-1].date()),
                'primary_output': 'forecast_annual.csv: real_bln_rub, area_mln_m2',
                'area_assumption': 'Segment base price and price dynamics equal the overall Moscow market.',
                'forecast_warning': 'No clipping: inspect negative forecasts and KS extrapolation flags.',
                'validation_warning': 'Known future KS; observed prices construct real actuals and auxiliary nominal validation. Overlapping 12-month horizons.'}
    (args.output / 'diagnostics.json').write_text(json.dumps(metadata, ensure_ascii=False,
                                                          indent=2, allow_nan=False), encoding='utf-8')
    if not args.no_charts: make_charts(data, train, model, history, args.output)
    print(f'Huber {args.scope}: real = {model["coef"][0]:.6f} + {model["coef"][1]:.6f} * KS_avg3')
    print('Annual result: real_bln_rub in January-2023 prices; area_mln_m2 in million m²')
    print(annual[['year', 'real_bln_rub', 'area_mln_m2', 'status']].to_string(index=False, float_format=lambda v: f'{v:.6f}'))
    print(f'Results: {args.output.resolve()}')


if __name__ == '__main__':
    main()
